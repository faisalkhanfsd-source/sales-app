"""Raw Data + File Format  ->  filled File Format (Sale & Coverage)."""
import io
import pandas as pd
import openpyxl

SALE_COL = "Sale Case Conv 250ml 2"
GROUPS = {"CSD", "NCB", "LRB"}


def _norm(s):
    return str(s).strip().upper() if pd.notna(s) else ""


def load_raw(file_like):
    df = pd.read_excel(file_like, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    need = [SALE_COL, "OUTLET CODE", "Brand", "Flavor", "Pack3",
            "PEPSI TRADE CHANNELS", "CSD", "NCB", "LRB"]
    missing = [c for c in need if c not in df.columns]
    if missing:
        raise ValueError("Raw Data میں یہ کالم نہیں ملے: " + ", ".join(missing))
    # repeated header rows inside the data -> sale is not a number -> dropped
    df["_sale"] = pd.to_numeric(df[SALE_COL], errors="coerce")
    df = df[df["_sale"].notna()].copy()
    df["_chan"] = df["PEPSI TRADE CHANNELS"].map(_norm)
    df["_brand"] = df["Brand"].map(_norm)
    df["_flavor"] = df["Flavor"].map(_norm)
    df["_pack"] = df["Pack3"].map(_norm)
    df["_outlet"] = df["OUTLET CODE"].astype(str).str.strip()
    for g in GROUPS:
        df["_f_" + g] = df[g].map(_norm).eq("Y")
    return df


def _calc(sub, row_brand, row_flavor, row_pack):
    """Return (sale, coverage) for one line of the format."""
    b, f, p = _norm(row_brand), _norm(row_flavor), _norm(row_pack)
    if b in GROUPS:
        sub = sub[sub["_f_" + b]]
    else:
        sub = sub[sub["_brand"] == b]
        if p == "TOTAL":
            pass                                  # whole brand (all flavors/packs)
        elif f:
            sub = sub[sub["_flavor"] == f]
    if p and p != "TOTAL":
        sub = sub[sub["_pack"] == p]
    if sub.empty:
        return 0.0, 0
    per_outlet = sub.groupby("_outlet")["_sale"].sum()
    return float(sub["_sale"].sum()), int((per_outlet > 0).sum())


def build(raw_file, format_file):
    raw = load_raw(raw_file)
    wb = openpyxl.load_workbook(format_file)
    ws = wb.worksheets[0]
    hdr = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}
    cc, cb, cf, cp = hdr["channel"], hdr["brand"], hdr["flavor"], hdr["pack"]
    cs, cv = hdr["sale"], hdr["coverage"]

    by_chan = {k: v for k, v in raw.groupby("_chan")}
    cache, n = {}, 0
    for r in range(2, ws.max_row + 1):
        ch = ws.cell(r, cc).value
        if ch is None:
            continue
        brand, flav, pack = (ws.cell(r, c).value for c in (cb, cf, cp))
        key = (_norm(ch), _norm(brand), _norm(flav), _norm(pack))
        if key not in cache:
            sub = raw if key[0] == "OVERALL" else by_chan.get(key[0], raw.iloc[0:0])
            cache[key] = _calc(sub, brand, flav, pack)
        sale, cov = cache[key]
        ws.cell(r, cs).value = round(sale, 4)
        ws.cell(r, cv).value = cov
        n += 1
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    stats = {"rows": n, "raw_rows": len(raw), "raw_total": float(raw["_sale"].sum()),
             "unmapped_rows": int((~raw["_chan"].isin(
                 {"CONVENTIONAL TRADE", "HORECA", "SPECIALITY STORE", "MODERN TRADE",
                  "WHOLESALE", "INSTITUTIONS"})).sum())}
    return out, stats
