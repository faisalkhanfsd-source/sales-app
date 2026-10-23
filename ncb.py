"""SALES_VS_TARGET workbook -> PNG of the SUMMARY sheet (one region block), Excel-like look."""
import io, os, re, base64, json, urllib.request
import openpyxl
from PIL import Image, ImageDraw, ImageFont

COLS = (1, 2, 4, 5, 6, 7, 8, 9)            # A, B(+C merged), D..I
WIDTHS = [120, 250, 100, 100, 72, 100, 84, 72]
THEME = {(0, 0): "#FFFFFF", (1, 0): "#000000", (4, 0.4): "#8EA9DB"}
FONT_DIRS = ["/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf",
             "/usr/share/fonts/dejavu/DejaVuSans%s.ttf"]


def _font(size, bold=False):
    for p in FONT_DIRS:
        f = p % ("-Bold" if bold else "")
        if os.path.exists(f):
            return ImageFont.truetype(f, size)
    try:
        return ImageFont.load_default(size)
    except TypeError:
        return ImageFont.load_default()


def _hex(color, default=None):
    if color is None:
        return default
    try:
        if color.type == "rgb" and isinstance(color.rgb, str):
            return "#" + color.rgb[-6:]
        if color.type == "theme":
            return THEME.get((color.theme, round(color.tint or 0, 1)), default)
    except Exception:
        pass
    return default


def _style(cell):
    fill = _hex(cell.fill.fgColor) if cell.fill and cell.fill.fill_type else None
    col = _hex(cell.font.color, "#000000") if cell.font else "#000000"
    return {"fill": fill, "bold": bool(cell.font and cell.font.b), "color": col}


def read_summary(file_like):
    wb = openpyxl.load_workbook(file_like, data_only=True)
    ws = next((s for s in wb.worksheets if s.title.strip().upper() == "SUMMARY"), None)
    if ws is None:
        raise ValueError("فائل میں 'SUMMARY' نام کی شیٹ نہیں ملی")
    hdr = next((r for r in range(1, 15) if str(ws.cell(r, 1).value or "").strip().upper() == "REGION"), None)
    if hdr is None:
        raise ValueError("SUMMARY شیٹ میں REGION والی ہیڈر لائن نہیں ملی")
    date_txt = next((str(ws.cell(r, 1).value) for r in range(1, hdr)
                     if str(ws.cell(r, 1).value or "").startswith("From")), "")
    title_r = next((r for r in range(1, hdr) if ws.cell(r, 1).value
                    and not str(ws.cell(r, 1).value).startswith("From")), None)
    title = str(ws.cell(title_r, 1).value) if title_r else "MTD SALE Vs Target"
    sty = lambda r, c: _style(ws.cell(r, c))
    meta = {"date_style": sty(1 if not date_txt else next(r for r in range(1, hdr)
            if str(ws.cell(r, 1).value or "").startswith("From")), 1),
            "title_style": sty(title_r or hdr, 1), "mix_style": sty(title_r or hdr, 8),
            "head_style": [sty(hdr, c) for c in COLS]}
    heads = [ws.cell(hdr, c).value for c in COLS]
    regions = {}
    for r in range(hdr + 1, ws.max_row + 1):
        reg = ws.cell(r, 1).value
        if not reg:
            continue
        key = str(reg).strip().upper()
        regions.setdefault(key, []).append({
            "v": [ws.cell(r, c).value for c in COLS],
            "s": [_style(ws.cell(r, c)) if c != 2 else _style(ws.cell(r, 2)) for c in COLS]})
    return {"date": date_txt, "title": title, "heads": heads, "regions": regions, **meta}


def _num(v):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return "" if v in (None, "") else str(v)
    return f"{v:,.0f}" if abs(v) >= .5 else "-"


def _pct(v):
    try:
        return f"{float(v) * 100:.0f}%"
    except (TypeError, ValueError):
        return ""


def _scale(vals):
    """Excel 3-colour scale: min red, 50th percentile yellow, max green."""
    lo, hi, c = "#F8696B", "#63BE7B", "#FFEB84"
    xs = sorted(v for v in vals if isinstance(v, (int, float)))
    if not xs:
        return lambda v: None
    mn, mx = xs[0], xs[-1]
    md = xs[len(xs) // 2] if len(xs) % 2 else (xs[len(xs) // 2 - 1] + xs[len(xs) // 2]) / 2
    rgb = lambda h: tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))
    mix = lambda a, b, t: "#%02X%02X%02X" % tuple(round(x + (y - x) * t) for x, y in zip(rgb(a), rgb(b)))

    def f(v):
        if not isinstance(v, (int, float)):
            return None
        if v <= md:
            return mix(lo, c, 0 if md == mn else (v - mn) / (md - mn))
        return mix(c, hi, 0 if mx == md else (v - md) / (mx - md))
    return f


def caption(data, region):
    rows = data["regions"].get(region.strip().upper())
    if not rows:
        return region
    tot = rows[-1]["v"]
    return f"{region.strip().upper()} {_pct(tot[4])}"


def render(data, region="SBL TOTAL"):
    rows = data["regions"].get(region.strip().upper())
    if not rows:
        raise ValueError("ریجن نہیں ملا: " + region)
    H, pad = 30, 7
    W = sum(WIDTHS)
    FOOT = 46
    img = Image.new("RGB", (W + 1, H * (len(rows) + 3) + 1 + FOOT), "white")
    d = ImageDraw.Draw(img)
    f, fb = _font(14), _font(14, True)

    def cell(x, y, w, text, st, align="left", fill=None, color=None):
        d.rectangle([x, y, x + w, y + H], fill=fill or st.get("fill") or "white", outline="#8c8c8c")
        ft = fb if st.get("bold") else f
        tw = d.textlength(text, font=ft)
        tx = x + (w - tw) / 2 if align == "center" else x + w - tw - pad if align == "right" else x + pad
        d.text((tx, y + H / 2), text, font=ft, fill=color or st.get("color") or "black", anchor="lm")

    cell(0, 0, W, data["date"], {**data["date_style"], "bold": True}, "center")
    left = sum(WIDTHS[:6])
    cell(0, H, left, data["title"], {**data["title_style"], "bold": True}, "center")
    cell(left, H, W - left, "Target Pack Mix", {**data["mix_style"], "bold": True}, "center")
    x = 0
    for i, (w, h) in enumerate(zip(WIDTHS, data["heads"])):
        cell(x, 2 * H, w, str(h), {**data["head_style"][i], "bold": True}, "center" if i == 1 else "left")
        x += w
    is_tot = lambda r: str(r["v"][1] or "").strip().upper() in ("", "GTOTAL")
    detail = [r for r in rows if not is_tot(r)]
    sc = _scale([r["v"][7] for r in detail])
    for i, r in enumerate(rows):
        y = (3 + i) * H
        v = r["v"]
        total = is_tot(r)
        texts = [str(v[0] or ""), str(v[1] or ""), _num(v[2]), _num(v[3]), _pct(v[4]),
                 _num(v[5]), _pct(v[6]), _pct(v[7])]
        x = 0
        for j, (w, t) in enumerate(zip(WIDTHS, texts)):
            st, fill, color = r["s"][j], None, None
            if j == 7 and not total and sc(v[7]):
                fill, color = sc(v[7]), "#000000"
            cell(x, y, w, t, st, "center" if j == 1 and not total else ("left" if j < 2 else "right"),
                 fill, color)
            x += w
    fy = H * (len(rows) + 3) + 1
    d.rectangle([0, fy, W, fy + FOOT], fill="#FFFFFF")
    d.text((14, fy + FOOT / 2), caption(data, region), font=_font(26, True), fill="#002060", anchor="lm")
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


def norm_phone(p):
    d = re.sub(r"\D", "", str(p or ""))
    if d.startswith("00"):
        d = d[2:]
    if d.startswith("0"):
        d = "92" + d[1:]
    elif len(d) == 10 and d.startswith("3"):
        d = "92" + d
    return d


def read_phones(file_like, filename=""):
    """Phone list file (xlsx or csv) with columns PHONE NO, REGION -> [(phone, REGION)]."""
    if filename.lower().endswith(".csv"):
        import csv
        txt = file_like.read().decode("utf-8-sig", "ignore").splitlines()
        rows = list(csv.reader(txt))
    else:
        ws = openpyxl.load_workbook(file_like, data_only=True).worksheets[0]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError("نمبروں والی فائل خالی ہے")
    head = [str(c or "").strip().upper() for c in rows[0]]
    try:
        ip, ir = next(i for i, h in enumerate(head) if "PHONE" in h), head.index("REGION")
    except (StopIteration, ValueError):
        raise ValueError("نمبروں والی فائل میں PHONE NO اور REGION کے کالم ہونے چاہییں")
    return [(norm_phone(r[ip]), str(r[ir]).strip().upper()) for r in rows[1:]
            if len(r) > max(ip, ir) and r[ip] and r[ir]]


def send_whatsapp(png, to=None, text=None):
    """Optional auto-send through WhatsApp Cloud API (needs WA_TOKEN, WA_PHONE_ID)."""
    tok, pid = os.environ.get("WA_TOKEN"), os.environ.get("WA_PHONE_ID")
    to = to or os.environ.get("WA_TO")
    if not (tok and pid and to):
        return None
    b = "----wa" + base64.b16encode(os.urandom(6)).decode()
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"messaging_product\"\r\n\r\nwhatsapp\r\n"
            f"--{b}\r\nContent-Disposition: form-data; name=\"type\"\r\n\r\nimage/png\r\n"
            f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"report.png\"\r\n"
            f"Content-Type: image/png\r\n\r\n").encode() + png + f"\r\n--{b}--\r\n".encode()
    req = urllib.request.Request(f"https://graph.facebook.com/v20.0/{pid}/media", body,
                                 {"Authorization": "Bearer " + tok,
                                  "Content-Type": "multipart/form-data; boundary=" + b})
    mid = json.load(urllib.request.urlopen(req, timeout=60))["id"]
    img = {"id": mid}
    if text:
        img["caption"] = text
    msg = json.dumps({"messaging_product": "whatsapp", "to": to, "type": "image", "image": img}).encode()
    req = urllib.request.Request(f"https://graph.facebook.com/v20.0/{pid}/messages", msg,
                                 {"Authorization": "Bearer " + tok, "Content-Type": "application/json"})
    urllib.request.urlopen(req, timeout=60).read()
    return True
